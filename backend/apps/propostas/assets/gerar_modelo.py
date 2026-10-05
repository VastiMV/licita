"""Gera `apps/propostas/modelo_padrao.docx` — a proposta comercial timbrada.

O texto segue o modelo que a equipe usava (docs/MODELO DE DOCUMENTO
PROPOSTA.docx); o timbrado é o logo da Inside Solutions (`timbrado.jpg`,
recortado do mesmo Word). Os campos `{{ }}` e `{%tr %}` são do docxtpl —
ver a lista em `apps/propostas/documento.py`.

    python apps/propostas/assets/gerar_modelo.py

O .docx resultante é versionado (o build da imagem não roda este script).
Quem quiser mudar o texto baixa o modelo na tela, edita no Word e envia de
volta — não precisa deste script.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

AQUI = Path(__file__).parent
DESTINO = AQUI.parent / "modelo_padrao.docx"

NAVY = RGBColor(0x02, 0x0E, 0x26)
AZUL = RGBColor(0x1E, 0x8B, 0xF0)
CINZA = RGBColor(0x5B, 0x67, 0x7A)
TEXTO = RGBColor(0x1F, 0x29, 0x37)
FUNDO_NAVY = "020E26"
FUNDO_ZEBRA = "F2F6FB"
FUNDO_TOTAL = "E6F0FC"

LARGURA_PAGINA = Cm(21)
MARGEM = Cm(2)


def faixa_do_timbrado() -> Path:
    """O logo numa faixa navy da largura da página — o fundo do logo já é
    navy, então a faixa continua o fundo e o logo fica à esquerda."""

    logo = Image.open(AQUI / "timbrado.jpg").convert("RGB")
    altura = 300
    logo = logo.resize((round(logo.width * altura / logo.height), altura), Image.LANCZOS)
    faixa = Image.new("RGB", (2480, altura + 40), logo.getpixel((3, 3)))
    faixa.paste(logo, (150, 20))
    # Filete azul embaixo, a cor de "licitações públicas" no logo.
    for y in range(faixa.height - 8, faixa.height):
        for x in range(faixa.width):
            faixa.putpixel((x, y), (30, 139, 240))
    destino = AQUI / "_faixa.png"
    faixa.save(destino)
    return destino


def sombrear(celula, cor: str) -> None:
    propriedades = celula._tc.get_or_add_tcPr()
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), cor)
    propriedades.append(sombra)


def bordas(tabela, cor: str = "D5DEEA") -> None:
    """Só linhas horizontais finas: tabela de proposta, não planilha."""

    propriedades = tabela._tbl.tblPr
    elemento = OxmlElement("w:tblBorders")
    for lado in ("top", "bottom", "insideH"):
        borda = OxmlElement(f"w:{lado}")
        borda.set(qn("w:val"), "single")
        borda.set(qn("w:sz"), "4")
        borda.set(qn("w:color"), cor)
        elemento.append(borda)
    for lado in ("left", "right", "insideV"):
        borda = OxmlElement(f"w:{lado}")
        borda.set(qn("w:val"), "nil")
        elemento.append(borda)
    propriedades.append(elemento)


def repetir_cabecalho(linha) -> None:
    propriedades = linha._tr.get_or_add_trPr()
    elemento = OxmlElement("w:tblHeader")
    elemento.set(qn("w:val"), "true")
    propriedades.append(elemento)


def filete(paragrafo, cor: str = "1E8BF0", tamanho: str = "6") -> None:
    propriedades = paragrafo._p.get_or_add_pPr()
    borda = OxmlElement("w:pBdr")
    baixo = OxmlElement("w:top")
    baixo.set(qn("w:val"), "single")
    baixo.set(qn("w:sz"), tamanho)
    baixo.set(qn("w:space"), "6")
    baixo.set(qn("w:color"), cor)
    borda.append(baixo)
    propriedades.append(borda)


def texto(paragrafo, conteudo: str, *, negrito=False, cor=TEXTO, tamanho=None, italico=False):
    run = paragrafo.add_run(conteudo)
    run.bold = negrito
    run.italic = italico
    run.font.color.rgb = cor
    if tamanho:
        run.font.size = Pt(tamanho)
    return run


def paragrafo(doc, conteudo: str = "", *, depois=6, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY, **estilo):
    p = doc.add_paragraph()
    p.alignment = alinhamento
    p.paragraph_format.space_after = Pt(depois)
    if conteudo:
        texto(p, conteudo, **estilo)
    return p


def secao(doc, numero: str, titulo: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    texto(p, f"{numero}  ", negrito=True, cor=AZUL, tamanho=10.5)
    texto(p, titulo.upper(), negrito=True, cor=NAVY, tamanho=10.5)
    return p


def gerar() -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.font.color.rgb = TEXTO
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    secao_doc = doc.sections[0]
    secao_doc.page_width, secao_doc.page_height = LARGURA_PAGINA, Cm(29.7)
    secao_doc.left_margin = secao_doc.right_margin = MARGEM
    secao_doc.top_margin = Cm(4.2)
    secao_doc.bottom_margin = Cm(2.6)
    secao_doc.header_distance = Cm(0)
    secao_doc.footer_distance = Cm(0.8)

    # ---------- timbrado: a faixa sangra até as bordas ----------
    cabecalho = secao_doc.header.paragraphs[0]
    cabecalho.paragraph_format.left_indent = -MARGEM
    cabecalho.paragraph_format.right_indent = -MARGEM
    cabecalho.add_run().add_picture(str(faixa_do_timbrado()), width=LARGURA_PAGINA)

    rodape = secao_doc.footer.paragraphs[0]
    rodape.alignment = WD_ALIGN_PARAGRAPH.CENTER
    filete(rodape, "D5DEEA", "4")
    texto(rodape, "{{ empresa.nome }}", negrito=True, cor=NAVY, tamanho=8)
    texto(rodape, "  ·  CNPJ {{ empresa.cnpj }}", cor=CINZA, tamanho=8)
    linha2 = secao_doc.footer.add_paragraph()
    linha2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    texto(linha2, "{{ empresa.endereco }}{% if empresa.cep %} · CEP {{ empresa.cep }}{% endif %}", cor=CINZA, tamanho=8)
    linha3 = secao_doc.footer.add_paragraph()
    linha3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    texto(linha3, "{{ empresa.telefone }}  ·  {{ empresa.email }}", cor=AZUL, tamanho=8)

    # ---------- identificação ----------
    p = paragrafo(doc, depois=2, alinhamento=WD_ALIGN_PARAGRAPH.LEFT)
    texto(p, "PROPOSTA COMERCIAL", negrito=True, cor=NAVY, tamanho=18)
    p = paragrafo(doc, depois=12, alinhamento=WD_ALIGN_PARAGRAPH.LEFT)
    texto(p, "{{ edital.modalidade|upper }} Nº {{ edital.numero }}", negrito=True, cor=AZUL, tamanho=11)

    p = paragrafo(doc, depois=0, alinhamento=WD_ALIGN_PARAGRAPH.LEFT)
    texto(p, "À", cor=CINZA)
    p = paragrafo(doc, depois=0, alinhamento=WD_ALIGN_PARAGRAPH.LEFT)
    texto(p, "{{ orgao.nome }}", negrito=True)
    p = paragrafo(doc, depois=10, alinhamento=WD_ALIGN_PARAGRAPH.LEFT)
    texto(p, "UASG {{ orgao.uasg }}{% if orgao.cidade %} · {{ orgao.cidade }}{% endif %}", cor=CINZA)

    paragrafo(doc, "Prezados Senhores,", depois=8)

    secao(doc, "1.", "Apresentação")
    paragrafo(
        doc,
        "A empresa {{ empresa.nome }}, inscrita no CNPJ/MF sob o nº {{ empresa.cnpj }}, sediada "
        "na {{ empresa.endereco }}, por seu representante legal infra-assinado e qualificado, vem "
        "por meio desta apresentar sua proposta de preços para o {{ edital.modalidade }} nº "
        "{{ edital.numero }}, em conformidade com a sessão pública de lances ofertados por meio "
        "eletrônico{% if edital.prazo %} em {{ edital.prazo }}{% endif %}, DECLARANDO AINDA, sob "
        "as penas da Lei, ter pleno conhecimento do teor do respectivo edital, subordinando-se ao "
        "contido no mesmo, cujos dispositivos reconhece, para todos os efeitos, terem caráter "
        "contratual.",
    )

    secao(doc, "2.", "Cotação dos materiais, de acordo com o Termo de Referência")
    p = paragrafo(doc, depois=4, alinhamento=WD_ALIGN_PARAGRAPH.LEFT)
    texto(p, "Objeto: ", negrito=True, cor=NAVY, tamanho=9)
    texto(p, "{{ edital.objeto }}", cor=CINZA, tamanho=9)

    titulos = ["Item", "Descrição", "Marca", "Modelo", "Unid.", "Qtde", "Preço unit.", "Preço total"]
    larguras = [Cm(1.1), Cm(6.3), Cm(1.9), Cm(1.9), Cm(1.2), Cm(1.2), Cm(1.7), Cm(1.7)]
    numericas = {5, 6, 7}
    tabela = doc.add_table(rows=5, cols=len(titulos))
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabela.autofit = False
    bordas(tabela)

    for i, titulo in enumerate(titulos):
        celula = tabela.rows[0].cells[i]
        sombrear(celula, FUNDO_NAVY)
        celula.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        par = celula.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT if i in numericas else WD_ALIGN_PARAGRAPH.LEFT
        texto(par, titulo, negrito=True, cor=RGBColor(0xFF, 0xFF, 0xFF), tamanho=8.5)
    repetir_cabecalho(tabela.rows[0])

    tabela.rows[1].cells[0].paragraphs[0].add_run("{%tr for item in itens %}")
    campos = ["numero", "descricao", "marca", "modelo", "unidade", "quantidade", "valor_unitario", "valor_total"]
    for i, campo in enumerate(campos):
        celula = tabela.rows[2].cells[i]
        celula.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        par = celula.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT if i in numericas else WD_ALIGN_PARAGRAPH.LEFT
        texto(par, "{{ item.%s }}" % campo, negrito=campo == "valor_total", tamanho=8.5)
    tabela.rows[3].cells[0].paragraphs[0].add_run("{%tr endfor %}")

    total = tabela.rows[4].cells
    rotulo = total[0].merge(total[6])
    for celula in (rotulo, total[7]):
        sombrear(celula, FUNDO_TOTAL)
    rotulo.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    texto(rotulo.paragraphs[0], "VALOR GLOBAL", negrito=True, cor=NAVY, tamanho=9)
    total[7].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    texto(total[7].paragraphs[0], "{{ valor_total }}", negrito=True, cor=NAVY, tamanho=9)

    for linha in tabela.rows:
        for celula, largura in zip(linha.cells, larguras):
            celula.width = largura

    p = paragrafo(doc, depois=6, alinhamento=WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.space_before = Pt(8)
    texto(p, "Valor total: ", negrito=True, cor=NAVY)
    texto(p, "{{ valor_total }} ({{ valor_total_extenso }}).", negrito=True)

    secao(doc, "3.", "Validade da proposta")
    paragrafo(
        doc,
        "{{ validade_dias }} ({{ validade_extenso }}) dias, conforme edital e termo de referência, "
        "contados da data de sua apresentação.",
    )

    secao(doc, "4.", "Dados bancários para pagamento")
    p = paragrafo(doc, alinhamento=WD_ALIGN_PARAGRAPH.LEFT)
    texto(p, "Banco: ", negrito=True)
    texto(p, "{{ empresa.banco }}     ")
    texto(p, "Agência: ", negrito=True)
    texto(p, "{{ empresa.agencia }}     ")
    texto(p, "Conta-corrente: ", negrito=True)
    texto(p, "{{ empresa.conta }}")

    secao(doc, "5.", "Declarações")
    for declaracao in [
        "Declaramos que estamos de pleno acordo com todas as condições estabelecidas no Edital e "
        "seus Anexos, bem como aceitamos todas as obrigações e responsabilidades constantes das "
        "especificações.",
        "Declaramos que nos preços cotados estão incluídas todas as despesas que, direta ou "
        "indiretamente, fazem parte do presente objeto, tais como gastos da empresa com suporte "
        "técnico e administrativo, impostos, seguros, taxas, ou quaisquer outros que possam "
        "incidir sobre gastos da empresa, sem quaisquer acréscimos em virtude de expectativa "
        "inflacionária e deduzidos os descontos eventualmente concedidos.",
        "Declaramos ainda que a documentação de habilitação consiste nos documentos elencados no "
        "SICAF, conforme a Lei de Licitações nº 14.133/2021 e o competente Termo de Referência.",
        "Declaramos estar de acordo e que atendemos ao artigo 7º, inciso XXXIII, da Constituição "
        "Federal.",
    ]:
        p = paragrafo(doc, declaracao, depois=4)
        p.paragraph_format.left_indent = Cm(0.4)

    secao(doc, "6.", "Dados da empresa")
    dados = doc.add_table(rows=1, cols=2)
    dados.autofit = False
    esquerda, direita = dados.rows[0].cells
    esquerda.width = direita.width = Cm(8.5)
    sombrear(esquerda, FUNDO_ZEBRA)
    sombrear(direita, FUNDO_ZEBRA)
    for linha, negrito in [
        ("{{ empresa.nome }}", True),
        ("CNPJ {{ empresa.cnpj }}", False),
        ("{{ empresa.endereco }}", False),
        ("{% if empresa.cep %}CEP {{ empresa.cep }}{% endif %}", False),
    ]:
        par = esquerda.paragraphs[0] if not esquerda.paragraphs[0].text else esquerda.add_paragraph()
        texto(par, linha, negrito=negrito, cor=NAVY if negrito else TEXTO, tamanho=9)
    for linha in ["E-mail: {{ empresa.email }}", "Telefone: {{ empresa.telefone }}"]:
        par = direita.paragraphs[0] if not direita.paragraphs[0].text else direita.add_paragraph()
        texto(par, linha, tamanho=9)

    secao(doc, "7.", "Representante legal")
    paragrafo(
        doc,
        "{{ empresa.responsavel_legal|upper }}{% if empresa.responsavel_qualificacao %}, "
        "{{ empresa.responsavel_qualificacao }}{% endif %}{% if empresa.responsavel_cpf %}, "
        "portador do CPF nº {{ empresa.responsavel_cpf }}{% endif %}{% if empresa.responsavel_rg %} "
        "e do RG nº {{ empresa.responsavel_rg }}{% endif %}.",
    )

    p = paragrafo(doc, "{{ cidade_data }}", alinhamento=WD_ALIGN_PARAGRAPH.RIGHT, depois=36)
    p.paragraph_format.space_before = Pt(14)

    for conteudo, estilo in [
        ("______________________________________________", {"cor": CINZA}),
        ("{{ empresa.responsavel_legal|upper }}", {"negrito": True, "cor": NAVY}),
        ("{% if empresa.responsavel_cpf %}CPF {{ empresa.responsavel_cpf }} · {% endif %}Representante legal", {"cor": CINZA, "tamanho": 9}),
        ("{{ empresa.nome }}", {"cor": CINZA, "tamanho": 9}),
    ]:
        p = paragrafo(doc, depois=0, alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
        p.paragraph_format.keep_with_next = True
        texto(p, conteudo, **estilo)

    doc.save(DESTINO)
    (AQUI / "_faixa.png").unlink()


if __name__ == "__main__":
    gerar()
