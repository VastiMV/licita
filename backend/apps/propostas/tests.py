"""Propostas: gerar move a licitação do Cotador para a etapa Proposta (e,
com proposta, o prazo vencer não a leva para Encerradas); o módulo lista,
escolhe a empresa, gera o Word do modelo e guarda os arquivos."""

from __future__ import annotations

import datetime as dt
import io
import tempfile

from django.core.files.uploadedfile import SimpleUploadedFile
from docx import Document
from rest_framework.test import APITestCase

from apps.armazenamento.models import ConfigArmazenamento
from apps.cotador.test_views import item_cotado
from apps.empresas.models import Empresa
from apps.licitacoes import test_ciclo
from apps.licitacoes.models import EventoOportunidadeSalva
from apps.tenants.atual import tenant_atual

from .documento import MODELO_PADRAO
from .models import ArquivoProposta, Proposta


class GerarPropostaTests(APITestCase):
    setUp = test_ciclo.CicloTests.setUp
    salva = test_ciclo.CicloTests.salva
    cotar = test_ciclo.CicloTests.cotar
    quadro = test_ciclo.CicloTests.quadro
    coluna = test_ciclo.CicloTests.coluna

    def gerar(self, salva):
        return self.client.post("/api/propostas/", {"oportunidade_id": salva.pk}, format="json")

    def test_gerar_tira_do_cotador_e_poe_na_coluna_proposta(self):
        salva = self.salva("1")
        self.cotar(salva, itens=[item_cotado(valor_referencia="50.00")])

        resposta = self.gerar(salva)

        self.assertEqual(resposta.status_code, 201)
        self.assertEqual(self.client.get("/api/licitacoes/cotacoes/").data["results"], [])
        quadro = self.quadro()
        self.assertEqual(self.coluna(quadro, "cotacao")["cartoes"], [])
        coluna = self.coluna(quadro, "proposta")
        self.assertTrue(coluna["disponivel"])
        self.assertEqual([c["id"] for c in coluna["cartoes"]], [salva.pk])
        self.assertTrue(
            salva.eventos.filter(tipo=EventoOportunidadeSalva.Tipo.PROPOSTA).exists()
        )

    def test_gerar_de_novo_e_idempotente(self):
        salva = self.salva("1")
        self.cotar(salva)
        self.gerar(salva)

        self.assertEqual(self.gerar(salva).status_code, 200)
        self.assertEqual(Proposta.objects.count(), 1)
        self.assertEqual(salva.eventos.filter(tipo=EventoOportunidadeSalva.Tipo.PROPOSTA).count(), 1)

    def test_sem_cotacao_nao_gera(self):
        self.assertEqual(self.gerar(self.salva("1")).status_code, 400)

    def test_com_proposta_o_prazo_vencido_nao_encerra(self):
        salva = self.salva("1")
        self.cotar(salva)
        self.gerar(salva)
        salva.data_encerramento_proposta = self.hoje - dt.timedelta(days=1)
        salva.save()

        self.assertEqual(self.client.get("/api/licitacoes/encerradas/").data["count"], 0)
        self.assertEqual(
            [c["id"] for c in self.coluna(self.quadro(), "proposta")["cartoes"]], [salva.pk]
        )

    def test_cotacao_com_proposta_nao_pode_ser_apagada(self):
        salva = self.salva("1")
        self.cotar(salva)
        self.gerar(salva)
        salva.refresh_from_db()

        resposta = self.client.delete(f"/api/cotador/cotacoes/{salva.cotacao.pk}/")

        self.assertEqual(resposta.status_code, 409)


# ---------- o módulo: lista, empresa, Word e arquivos ----------


def _texto(conteudo: bytes) -> str:
    documento = Document(io.BytesIO(conteudo))
    partes = [p.text for p in documento.paragraphs]
    for tabela in documento.tables:
        for linha in tabela.rows:
            partes.extend(c.text for c in linha.cells)
    return "\n".join(partes)


class ModuloPropostasTests(APITestCase):
    setUp_ciclo = test_ciclo.CicloTests.setUp
    salva = test_ciclo.CicloTests.salva
    cotar = test_ciclo.CicloTests.cotar

    def setUp(self):
        self.setUp_ciclo()
        self.tenant = tenant_atual()
        self.empresa = Empresa.objects.create(
            tenant=self.tenant,
            nome="Inside Solutions Ltda",
            cnpj="11222333000181",
            cidade="Curitiba",
            uf="PR",
            responsavel_legal="Vasti Marucci",
        )
        ConfigArmazenamento.objects.create(
            tenant=self.tenant, driver="local", opcoes={"raiz": tempfile.mkdtemp()}
        )
        self.oportunidade = self.salva("1")
        self.cotar(self.oportunidade, itens=[item_cotado(valor_referencia="50.00")])
        resposta = self.client.post(
            "/api/propostas/", {"oportunidade_id": self.oportunidade.pk}, format="json"
        )
        self.proposta_id = resposta.data["id"]

    def test_com_uma_empresa_so_a_proposta_ja_nasce_com_ela(self):
        self.assertEqual(
            self.client.get(f"/api/propostas/{self.proposta_id}/").data["empresa"]["id"],
            self.empresa.pk,
        )

    def test_lista_so_as_que_estao_em_proposta(self):
        self.cotar(self.salva("2"))  # em cotação: não entra

        linhas = self.client.get("/api/propostas/").data["results"]

        self.assertEqual([l["id"] for l in linhas], [self.oportunidade.pk])
        self.assertEqual(linhas[0]["proposta"]["id"], self.proposta_id)
        self.assertFalse(linhas[0]["proposta"]["documento_gerado"])

    def test_gera_o_word_do_modelo_padrao_e_guarda_entre_os_arquivos(self):
        resposta = self.client.get(f"/api/propostas/{self.proposta_id}/documento/")

        self.assertEqual(resposta.status_code, 200)
        texto = _texto(resposta.content)
        self.assertIn("Inside Solutions Ltda", texto)
        self.assertIn("11.222.333/0001-81", texto)
        self.assertIn("Edital 1", texto)
        self.assertIn("Vasti Marucci", texto)
        self.assertNotIn("{{", texto)
        self.assertEqual(
            ArquivoProposta.objects.filter(origem=ArquivoProposta.Origem.GERADO).count(), 1
        )

        # Gerar de novo substitui, não acumula.
        self.client.get(f"/api/propostas/{self.proposta_id}/documento/")
        self.assertEqual(
            ArquivoProposta.objects.filter(origem=ArquivoProposta.Origem.GERADO).count(), 1
        )

    def test_modelo_do_cliente_substitui_o_padrao(self):
        modelo = Document()
        modelo.add_paragraph("MODELO DA VASTI {{ empresa.nome }} {{ valor_total }}")
        conteudo = io.BytesIO()
        modelo.save(conteudo)
        upload = SimpleUploadedFile("modelo.docx", conteudo.getvalue())

        self.assertEqual(
            self.client.post("/api/propostas/modelo/", {"arquivo": upload}).status_code, 201
        )
        texto = _texto(self.client.get(f"/api/propostas/{self.proposta_id}/documento/").content)

        self.assertIn("MODELO DA VASTI Inside Solutions Ltda R$", texto)

        self.client.delete("/api/propostas/modelo/")
        self.assertTrue(self.client.get("/api/propostas/modelo/").data["padrao"])

    def test_modelo_que_nao_e_word_e_recusado(self):
        upload = SimpleUploadedFile("modelo.pdf", b"%PDF")
        self.assertEqual(
            self.client.post("/api/propostas/modelo/", {"arquivo": upload}).status_code, 400
        )

    def test_sobe_baixa_e_remove_arquivo_da_licitacao(self):
        upload = SimpleUploadedFile("declaracao.pdf", b"%PDF-1.4 x", content_type="application/pdf")

        criado = self.client.post(f"/api/propostas/{self.proposta_id}/arquivos/", {"arquivo": upload})

        self.assertEqual(criado.status_code, 201)
        self.assertEqual(criado.data["nome"], "declaracao.pdf")
        baixar = self.client.get(f"/api/propostas/arquivos/{criado.data['id']}/download/")
        self.assertEqual(baixar.status_code, 200)
        self.assertEqual(
            self.client.delete(f"/api/propostas/arquivos/{criado.data['id']}/").status_code, 204
        )
        self.assertEqual(self.client.get(f"/api/propostas/{self.proposta_id}/").data["arquivos"], [])

    def test_trocar_a_empresa(self):
        outra = Empresa.objects.create(tenant=self.tenant, nome="Filial", cnpj="11222333000262")

        resposta = self.client.patch(
            f"/api/propostas/{self.proposta_id}/", {"empresa_id": outra.pk}, format="json"
        )

        self.assertEqual(resposta.data["empresa"]["id"], outra.pk)

    def test_o_modelo_padrao_existe(self):
        self.assertTrue(MODELO_PADRAO.exists())
